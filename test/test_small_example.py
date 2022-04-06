"""Test the processing of a small dataset.

- Load a small dataset, in interleaved format.
- Add metadata, and split out a separate CognateTable and CognatesetTable.
- Add a LanguageTable and a ParameterTable for the concepts.
- Add segments
"""

import csv
import logging
import tempfile
from pathlib import Path

import pytest
import openpyxl

from lexedata import util
from lexedata.importer import excel_interleaved
from lexedata.edit import add_cognate_table
from lexedata.util.add_metadata import add_metadata
from lexedata.cli import logger


def assert_datasets_are_equal(ds1, ds2):
    """Check that two datasets are largely equal.

    Check that they contain the same table types, and that each table with a
    proper type contains the same rows, up to re-ordering.

    """
    tables1 = {ds1.get_tabletype(table) for table in ds1.tables}
    tables2 = {ds2.get_tabletype(table) for table in ds2.tables}
    assert (
        not tables1 ^ tables2
    ), f"Datasets contain different table types {tables1} and {tables2}."
    for table in ds1.tables:
        tabletype = ds1.get_tabletype(table)
        if tabletype is None:
            continue
        rows1 = [dict(row) for row in ds1[tabletype]]
        rows2 = [dict(row) for row in ds2[tabletype]]
        for row in rows1:
            assert (
                row in rows2
            ), f"Row {row} of {tabletype} 1 not found in {tabletype} 2"
        for row in rows2:
            assert (
                row in rows1
            ), f"Row {row} of {tabletype} 2 not found in {tabletype} 1"


@pytest.fixture
def interleaved_excel_example():
    data = [
        ["", "Duala", "Ntomba", "Ngombe", "Bushoong"],
        ["all", "ɓɛ́sɛ̃", "(nk)umá", "ńsò", "kim"],
        ["", 1, 9, 10, 11],
        ["arm", "dia", "lobɔ́kɔ", "lò-bókò (PL: màbókò)", "lɔ̀ɔ́"],
        ["", 7, 1, 1, 1],
        ["ashes", "mabúdú", "metókó", "búdùlù ~ pùdùlù", "bu-tók"],
        ["", 17, 16, 17, 16],
        ["bark", "bwelé", "lopoho ~ mpoho ~ lòpòhó", "émpósù ~ ímpósù", "yooʃ"],
        ["", 23, 22, 22, 22],
        ["belly", "dibum", "ikundú", "lì-bùmù", "ì-kù:n"],
        ["", 1, 18, 1, 18],
        ["big", "éndɛ̃nɛ̀", "nɛ́nɛ́", "nɛ́nɛ ~ nɛ́nɛ́nɛ", "nɛ́n"],
        ["", 1, 1, 1, 1],
        ["bird", "inɔ̌n", "mpulú", "é-mbùlù ~ í-mbùlù", "pul"],
        ["", 1, 7, 7, 7],
        ["bite", "kukwa", "lamata", "kokala", "a-ʃum"],
        ["", 6, 2, 7, 1],
        ["black", "wínda", "", "hínda; épííndu", "a-picy; ndwɛɛm"],
        ["", 21, "", "21, 21", "22, 23"],
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in data:
        ws.append(row)
    return ws


@pytest.fixture
def formtable_only_example():
    return util.fs.new_wordlist(
        FormTable=[
            {
                "ID": "duala_all",
                "Language_ID": "Duala",
                "Concept_ID": "all",
                "Form": "ɓɛ́sɛ̃",
                "Comment": None,
                "Cognateset_ID": "1",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "duala_arm",
                "Language_ID": "Duala",
                "Concept_ID": "arm",
                "Form": "dia",
                "Comment": None,
                "Cognateset_ID": "7",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "duala_ashes",
                "Language_ID": "Duala",
                "Concept_ID": "ashes",
                "Form": "mabúdú",
                "Comment": None,
                "Cognateset_ID": "17",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "duala_bark",
                "Language_ID": "Duala",
                "Concept_ID": "bark",
                "Form": "bwelé",
                "Comment": None,
                "Cognateset_ID": "23",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "duala_belly",
                "Language_ID": "Duala",
                "Concept_ID": "belly",
                "Form": "dibum",
                "Comment": None,
                "Cognateset_ID": "1",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "duala_big",
                "Language_ID": "Duala",
                "Concept_ID": "big",
                "Form": "éndɛ̃nɛ̀",
                "Comment": None,
                "Cognateset_ID": "1",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "duala_bird",
                "Language_ID": "Duala",
                "Concept_ID": "bird",
                "Form": "inɔ̌n",
                "Comment": None,
                "Cognateset_ID": "1",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "duala_bite",
                "Language_ID": "Duala",
                "Concept_ID": "bite",
                "Form": "kukwa",
                "Comment": None,
                "Cognateset_ID": "6",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "duala_black",
                "Language_ID": "Duala",
                "Concept_ID": "black",
                "Form": "wínda",
                "Comment": None,
                "Cognateset_ID": "21",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ntomba_all",
                "Language_ID": "Ntomba",
                "Concept_ID": "all",
                "Form": "(nk)umá",
                "Comment": None,
                "Cognateset_ID": "9",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ntomba_arm",
                "Language_ID": "Ntomba",
                "Concept_ID": "arm",
                "Form": "lobɔ́kɔ",
                "Comment": None,
                "Cognateset_ID": "1",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ntomba_ashes",
                "Language_ID": "Ntomba",
                "Concept_ID": "ashes",
                "Form": "metókó",
                "Comment": None,
                "Cognateset_ID": "16",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ntomba_bark",
                "Language_ID": "Ntomba",
                "Concept_ID": "bark",
                "Form": "lopoho ~ mpoho ~ lòpòhó",
                "Comment": None,
                "Cognateset_ID": "22",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ntomba_belly",
                "Language_ID": "Ntomba",
                "Concept_ID": "belly",
                "Form": "ikundú",
                "Comment": None,
                "Cognateset_ID": "18",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ntomba_big",
                "Language_ID": "Ntomba",
                "Concept_ID": "big",
                "Form": "nɛ́nɛ́",
                "Comment": None,
                "Cognateset_ID": "1",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ntomba_bird",
                "Language_ID": "Ntomba",
                "Concept_ID": "bird",
                "Form": "mpulú",
                "Comment": None,
                "Cognateset_ID": "7",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ntomba_bite",
                "Language_ID": "Ntomba",
                "Concept_ID": "bite",
                "Form": "lamata",
                "Comment": None,
                "Cognateset_ID": "2",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ngombe_all",
                "Language_ID": "Ngombe",
                "Concept_ID": "all",
                "Form": "ńsò",
                "Comment": None,
                "Cognateset_ID": "10",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ngombe_arm",
                "Language_ID": "Ngombe",
                "Concept_ID": "arm",
                "Form": "lò-bókò (PL: màbókò)",
                "Comment": None,
                "Cognateset_ID": "1",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ngombe_ashes",
                "Language_ID": "Ngombe",
                "Concept_ID": "ashes",
                "Form": "búdùlù ~ pùdùlù",
                "Comment": None,
                "Cognateset_ID": "17",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ngombe_bark",
                "Language_ID": "Ngombe",
                "Concept_ID": "bark",
                "Form": "émpósù ~ ímpósù",
                "Comment": None,
                "Cognateset_ID": "22",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ngombe_belly",
                "Language_ID": "Ngombe",
                "Concept_ID": "belly",
                "Form": "lì-bùmù",
                "Comment": None,
                "Cognateset_ID": "1",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ngombe_big",
                "Language_ID": "Ngombe",
                "Concept_ID": "big",
                "Form": "nɛ́nɛ ~ nɛ́nɛ́nɛ",
                "Comment": None,
                "Cognateset_ID": "1",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ngombe_bird",
                "Language_ID": "Ngombe",
                "Concept_ID": "bird",
                "Form": "é-mbùlù ~ í-mbùlù",
                "Comment": None,
                "Cognateset_ID": "7",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ngombe_bite",
                "Language_ID": "Ngombe",
                "Concept_ID": "bite",
                "Form": "kokala",
                "Comment": None,
                "Cognateset_ID": "7",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ngombe_black",
                "Language_ID": "Ngombe",
                "Concept_ID": "black",
                "Form": "hínda",
                "Comment": None,
                "Cognateset_ID": "21",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "ngombe_black_s2",
                "Language_ID": "Ngombe",
                "Concept_ID": "black",
                "Form": "épííndu",
                "Comment": None,
                "Cognateset_ID": "21",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "bushoong_all",
                "Language_ID": "Bushoong",
                "Concept_ID": "all",
                "Form": "kim",
                "Comment": None,
                "Cognateset_ID": "11",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "bushoong_arm",
                "Language_ID": "Bushoong",
                "Concept_ID": "arm",
                "Form": "lɔ̀ɔ́",
                "Comment": None,
                "Cognateset_ID": "1",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "bushoong_ashes",
                "Language_ID": "Bushoong",
                "Concept_ID": "ashes",
                "Form": "bu-tók",
                "Comment": None,
                "Cognateset_ID": "16",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "bushoong_bark",
                "Language_ID": "Bushoong",
                "Concept_ID": "bark",
                "Form": "yooʃ",
                "Comment": None,
                "Cognateset_ID": "22",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "bushoong_belly",
                "Language_ID": "Bushoong",
                "Concept_ID": "belly",
                "Form": "ì-kù:n",
                "Comment": None,
                "Cognateset_ID": "18",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "bushoong_big",
                "Language_ID": "Bushoong",
                "Concept_ID": "big",
                "Form": "nɛ́n",
                "Comment": None,
                "Cognateset_ID": "1",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "bushoong_bird",
                "Language_ID": "Bushoong",
                "Concept_ID": "bird",
                "Form": "pul",
                "Comment": None,
                "Cognateset_ID": "7",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "bushoong_bite",
                "Language_ID": "Bushoong",
                "Concept_ID": "bite",
                "Form": "a-ʃum",
                "Comment": None,
                "Cognateset_ID": "1",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "bushoong_black",
                "Language_ID": "Bushoong",
                "Concept_ID": "black",
                "Form": "a-picy",
                "Comment": None,
                "Cognateset_ID": "22",
                "Segments": [],
                "Source": [],
            },
            {
                "ID": "bushoong_black_s2",
                "Language_ID": "Bushoong",
                "Concept_ID": "black",
                "Form": "ndwɛɛm",
                "Comment": None,
                "Cognateset_ID": "23",
                "Segments": [],
                "Source": [],
            },
        ]
    )


def test_interleaved(interleaved_excel_example):
    ids = set()
    forms = [
        dict(
            zip(
                ["ID", "Language_ID", "Concept_ID", "Form", "Comment", "Cognateset_ID"],
                row,
            )
        )
        for row in excel_interleaved.import_interleaved(
            interleaved_excel_example, logger=logger, ids=ids
        )
    ]

    assert len(list(forms)) == 37
    assert ids == {
        "bushoong_all",
        "bushoong_arm",
        "bushoong_ashes",
        "bushoong_bark",
        "bushoong_belly",
        "bushoong_big",
        "bushoong_bite",
        "bushoong_bird",
        "bushoong_black",
        "bushoong_black_s2",
        "duala_all",
        "duala_arm",
        "duala_ashes",
        "duala_bark",
        "duala_belly",
        "duala_big",
        "duala_bird",
        "duala_bite",
        "duala_black",
        "ngombe_all",
        "ngombe_arm",
        "ngombe_ashes",
        "ngombe_bark",
        "ngombe_belly",
        "ngombe_big",
        "ngombe_bird",
        "ngombe_bite",
        "ngombe_black",
        "ngombe_black_s2",
        "ntomba_all",
        "ntomba_arm",
        "ntomba_ashes",
        "ntomba_bark",
        "ntomba_belly",
        "ntomba_big",
        "ntomba_bird",
        "ntomba_bite",
    }


def test_interleaved_excel_example_header_wrong(caplog):
    data = [
        ["Concept", "", "", "", ""],
        ["", "Duala", "Ntomba", "Ngombe", "Bushoong"],
        ["all", "ɓɛ́sɛ̃", "(nk)umá", "ńsò", "kim"],
        ["", 1, 9, 10, 11],
        ["arm", "dia", "lobɔ́kɔ", "lò-bókò (PL: màbókò)", "lɔ̀ɔ́"],
        ["", 7, 1, 1, 1],
        ["ashes", "mabúdú", "metókó", "búdùlù ~ pùdùlù", "bu-tók"],
        ["", 17, 16, 17, 16],
        ["bark", "bwelé", "lopoho ~ mpoho ~ lòpòhó", "émpósù ~ ímpósù", "yooʃ"],
        ["", 23, 22, 22, 22],
        ["belly", "dibum", "ikundú", "lì-bùmù", "ì-kù:n"],
        ["", 1, 18, 1, 18],
        ["big", "éndɛ̃nɛ̀", "nɛ́nɛ́", "nɛ́nɛ ~ nɛ́nɛ́nɛ", "nɛ́n"],
        ["", 1, 1, 1, 1],
        ["bird", "inɔ̌n", "mpulú", "é-mbùlù ~ í-mbùlù", "pul"],
        ["", 1, 7, 7, 7],
        ["bite", "kukwa", "lamata", "kokala", "a-ʃum"],
        ["", 6, 2, 7, 1],
        ["black", "wínda", "", "hínda; épííndu", "a-picy; ndwɛɛm"],
        ["", 21, "", "21, 21", "22, 23"],
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in data:
        ws.append(row)
    with caplog.at_level(logging.ERROR):
        with pytest.raises(SystemExit):
            for row in excel_interleaved.import_interleaved(ws, ids=set()):
                pass
    assert "expected one or more forms" in caplog.text


def test_create_metadata_valid(interleaved_excel_example):
    forms = [
        dict(
            zip(
                ["ID", "Language_ID", "Concept_ID", "Form", "Comment", "Cognateset_ID"],
                row,
            )
        )
        for row in excel_interleaved.import_interleaved(interleaved_excel_example)
    ]

    path = Path(tempfile.mkdtemp())
    with (path / "forms.csv").open(
        "w", encoding="utf-8", newline=""
    ) as form_table_file:
        writer = csv.DictWriter(
            form_table_file,
            fieldnames=[
                "ID",
                "Language_ID",
                "Concept_ID",
                "Form",
                "Comment",
                "Cognateset_ID",
            ],
        )
        writer.writeheader()
        writer.writerows(forms)
    ds = add_metadata(path / "forms.csv")
    ds.write_metadata(path / "Wordlist-metadata.json")

    assert {f.name for f in path.iterdir()} == {"forms.csv", "Wordlist-metadata.json"}
    assert len(ds.tables) == 1, "Expected a single table"
    assert [c.name for c in ds["FormTable"].tableSchema.columns] == [
        "ID",
        "Language_ID",
        "Concept_ID",
        "Form",
        "Comment",
        "Cognateset_ID",
        "Segments",
        "Source",
    ]
    assert ds.validate()


def test_create_metadata_correct(interleaved_excel_example, formtable_only_example):
    forms = [
        dict(
            zip(
                ["ID", "Language_ID", "Concept_ID", "Form", "Comment", "Cognateset_ID"],
                row,
            )
        )
        for row in excel_interleaved.import_interleaved(
            interleaved_excel_example, logger=logger
        )
    ]

    path = Path(tempfile.mkdtemp())
    with (path / "forms.csv").open(
        "w", encoding="utf-8", newline=""
    ) as form_table_file:
        writer = csv.DictWriter(
            form_table_file,
            fieldnames=[
                "ID",
                "Language_ID",
                "Concept_ID",
                "Form",
                "Comment",
                "Cognateset_ID",
            ],
        )
        writer.writeheader()
        writer.writerows(forms)
    ds = add_metadata(path / "forms.csv")
    ds.write_metadata(path / "Wordlist-metadata.json")

    # Normalize dataset
    ds.write(FormTable=list(ds["FormTable"]))
    assert_datasets_are_equal(ds, formtable_only_example)


def test_add_cog_tables_valid(formtable_only_example):
    ds = formtable_only_example
    add_cognate_table.add_cognate_table(ds, True)
    ds.validate()


def test_add_cog_tables_correct(formtable_only_example):
    ds = formtable_only_example
    add_cognate_table.add_cognate_table(ds, True)
    # TODO: Check whether the outcome is correct


# TODO: Add segments, add alignments, add language table, add concept table,
# add concepticon to concept table, find central concepts, phylogenetics
# exporter. Check some reports: In particular coverage. At a later stage, add a
# concept with homophones, and check they are found.
