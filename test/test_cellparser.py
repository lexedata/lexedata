# -*- coding: utf-8 -*-
import json
from pathlib import Path
from unittest import TestCase, TextTestRunner

import openpyxl as op

from lexedata.importer.cellparser import CellParser
from lexedata.database.objects import Form, Concept


class TestCellparser(TestCase):
    def setUp(self, file="test_cellparser.xlsx"):
        # set path and check for existing excel test file
        file = Path.cwd() / file
        if not file.exists():
            print("There are no test files provided")
            exit()

        wb = op.load_workbook(filename=file)
        sheets = wb.sheetnames
        wb = wb[sheets[0]]
        # for collecting different types of cells to be tested
        self.clean_cells = []  # all good
        self.tolerated_cells = []  # cellparser will change somethings
        self._error_cells = []  # error will be

        iter_concept = wb.iter_rows(min_row=3, max_col=6)  # iterates over rows with concepts
        iter_lan = wb.iter_cols(min_row=1, max_row=2, min_col=7)
        for row in wb.iter_rows(min_row=3, min_col=7):
            if row[0].row <= 5:
                for cell in row:
                    self.clean_cells.append(cell)
            if row[0].row == 6:
                for cell in row:
                    self.tolerated_cells.append(cell)
            if row[0].row > 6:
                for cell in row:
                    self._error_cells.append(cell)


def main():
    test = TestCellparser()
    test.setUp()
    empyt_concept = Concept(id="NA", set="NA", english="NA", english_strict="NA", spanish="NA", portuguese="NA",
                            french="NA", concept_comment="NA")
    mydump = {}
    for cell in test.clean_cells:
        forms = []
        for f in CellParser(cell):
            form = Form.create_form(f, "no_lan", cell, empyt_concept)
            form_dic = vars(form)
            del form_dic["_sa_instance_state"]
            forms.append(form_dic)
        mydump[cell.coordinate] = forms

    with (Path.cwd() / "clean_cells.json").open("w", encoding="utf8") as out:
        json.dump(mydump, out, indent=4, sort_keys=True)

def more_tests():
    c1 = Tester(
        "<tatatĩ>(humo){Guasch1962:717}, <timbo>(vapor, vaho, humareda, humo){Guasch1962:729};<tĩ> (humo, vapor de agua)$LDM:deleted 'nariz, pico, hocico, punta, and ápica' meanings; source incorrectly merges 'point' and 'smoke' meanings ${Guasch1962:729}")
    c2 = Tester(
        "/pãlĩ/ (de froment = wheat) (NCP: loan from french farine), /pɨlatɨ/ (de mais cru), /kuʔi/ (de mais grillé)")
    c3 = Tester(
        "/pãlĩ/ (de froment = wheat(NCP: loan from french farine), &<dummy>), /pɨlatɨ/ (de mais cru), /kuʔi/ (de mais grillé)")
    c4 = Tester("/popɨãpat/ 'back of elbow' {4}")
    c5 = Tester("<ayu> (nominalized &/afaaa/ version of &/ete/ 'about') {2}")
    for ele in [c1, c2, c3, c4, c5]:
        print(ele.value)
        print("is represented as: ")
        for f in CogCellParser(ele):
            print(f)
        input()
